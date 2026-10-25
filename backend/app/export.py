"""Содирот ба Word (.docx). PDF тавассути саҳифаи чоп дар frontend сохта мешавад."""

import io

from docx import Document
from docx.shared import Cm, Pt, RGBColor

LETTERS = "ABCDEF"
GRAY = RGBColor(0x5B, 0x6A, 0x72)


def _type_name(q) -> str:
    return getattr(q.type, "value", q.type)


def build_docx(test, with_answers: bool) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)

    title = doc.add_paragraph()
    run = title.add_run(test.title + (" - Answer key" if with_answers else ""))
    run.bold = True
    run.font.size = Pt(18)

    info = doc.add_paragraph()
    level = getattr(test.level, "value", test.level)
    run = info.add_run(f"{level}  |  {test.topic}")
    run.font.color.rgb = GRAY

    questions = [tq.question for tq in test.test_questions]
    if not with_answers:
        line = doc.add_paragraph()
        line.paragraph_format.space_before = Pt(8)
        line.add_run(f"Name: ______________________    Date: ____________    Score: ____ / {len(questions)}")

    for n, q in enumerate(questions, start=1):
        kind = _type_name(q)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{n}. ").bold = True
        p.add_run(q.prompt)

        if kind == "multiple_choice":
            for i, option in enumerate(q.options or []):
                op = doc.add_paragraph()
                op.paragraph_format.left_indent = Cm(0.9)
                op.paragraph_format.space_after = Pt(0)
                correct = with_answers and i == q.correct_answer
                r = op.add_run(f"{LETTERS[i]}) {option}" + ("  \u2713" if correct else ""))
                r.bold = correct
        elif kind == "true_false":
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Cm(0.9)
            for label, value in (("True", True), ("False", False)):
                r = op.add_run(label)
                if with_answers and q.correct_answer is value:
                    r.bold = True
                    r.underline = True
                if value:
                    op.add_run("  /  ")
        elif kind == "fill_blank" and with_answers:
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Cm(0.9)
            op.add_run("Answer: " + " / ".join(q.correct_answer)).bold = True

        if with_answers and q.explanation:
            ep = doc.add_paragraph()
            ep.paragraph_format.left_indent = Cm(0.9)
            r = ep.add_run(q.explanation)
            r.italic = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = GRAY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
